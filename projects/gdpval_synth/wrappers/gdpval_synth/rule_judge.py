"""Rule-based criterion verification for GDPVal Synth.

Can be used as a standalone script or imported as a module.

Standalone usage:
    TASK_WORKSPACE=/path/to/task \
    RUBRIC_PATH=/path/to/rubric.json \
    DELIVERABLE_SPEC_PATH=/path/to/spec.json \
    JUDGER_NAME=rule_judge \
    python rule_judge.py

Output: JudgerResult JSON to stdout.
"""
from __future__ import annotations

import json
import os
import re
import sys
from dataclasses import dataclass

import openpyxl
import pdfplumber
from docx import Document as DocxDocument
from pptx import Presentation


# ─── Expected field parsers ──────────────────────────────────────────────────

_FILENAME_RE = re.compile(r"[\w\-]+\.(?:xlsx|xls|xlsm|pdf|docx|doc|pptx|ppt)")


def parse_expected_filenames(expected: str) -> list[str]:
    """Extract filenames from expected text."""
    # Pattern 1: Files: ['a.xlsx', 'b.pdf']
    m = re.search(r"Files:\s*\[([^\]]+)\]", expected)
    if m:
        items = re.findall(r"'([^']+)'", m.group(1))
        return [i for i in items if "." in i]

    # Pattern 2: File named 'xxx.xlsx'
    m = re.search(r"[Ff]ile\s+named?\s+'([^']+)'", expected)
    if m:
        names = [m.group(1)]
        rest = expected[m.end():]
        more = _FILENAME_RE.findall(rest)
        names.extend(more)
        return list(dict.fromkeys(names))

    # Pattern 3: general filename scan
    all_fnames = _FILENAME_RE.findall(expected)
    if all_fnames:
        return list(dict.fromkeys(all_fnames))

    return []


def parse_expected_sheets(expected: str) -> list[str]:
    """Extract sheet names from expected text."""
    sheets = []
    # Pattern 1: sheets named 'X' (and 'Y' ...)
    for m in re.finditer(r"(?:sheet|sheets|worksheet|tab)s?\s+named\s+'([^']+)'((?:\s+and\s+'[^']+')*)", expected, re.IGNORECASE):
        sheets.append(m.group(1))
        if m.group(2):
            sheets.extend(re.findall(r"'([^']+)'", m.group(2)))
    # Pattern 2: 'X' and 'Y' sheets
    m = re.search(r"'([^']+)'(?:\s+and\s+'([^']+)')?\s+sheets", expected, re.IGNORECASE)
    if m:
        sheets.append(m.group(1))
        if m.group(2):
            sheets.append(m.group(2))
    # Pattern 3: A sheet named 'X'
    for m in re.finditer(r"[Aa]\s+sheet\s+named\s+'([^']+)'", expected):
        sheets.append(m.group(1))

    return list(dict.fromkeys(sheets))


def parse_expected_columns(expected: str) -> list[str]:
    """Extract column names from expected text."""
    _COL_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_ ]*[A-Za-z0-9_]$|^[A-Za-z_]$")

    # Pattern 1: [Excel ]columns[: order](X, Y, Z)
    m = re.search(r"(?:Excel\s+)?columns?(?:\s+order)?[\s:(\[]+([A-Za-z_][\w,\s]+?)(?:\)|;|\]|$)", expected)
    if m:
        raw = m.group(1)
        cols = [c.strip() for c in raw.split(",") if c.strip()]
        cols = [c for c in cols if _COL_RE.match(c)]
        if cols:
            return cols

    # Pattern 2: contains columns: X, Y, Z
    m = re.search(r"contains?\s+columns?[\s:]+([A-Za-z_][\w,\s]+?)(?:\.|;|$)", expected)
    if m:
        raw = m.group(1)
        cols = [c.strip() for c in raw.split(",") if c.strip()]
        cols = [c for c in cols if _COL_RE.match(c)]
        if cols:
            return cols

    return []


def parse_expected_sections(expected: str) -> list[str]:
    """Extract document section names from expected text."""
    # Pattern 1: PDF Sections: ['X', 'Y', 'Z']
    m = re.search(r"[Ss]ections?:\s*\[([^\]]+)\]", expected)
    if m:
        return [s.strip().strip("'\"") for s in m.group(1).split(",") if s.strip()]

    # Pattern 2: sections: X, Y, Z
    m = re.search(r"(?:following\s+)?(?:\w+\s+)?sections?:\s*(.+?)(?:\.|;|$)", expected, re.IGNORECASE)
    if m:
        raw = m.group(1).strip()
        parts = re.split(r",\s*(?:and\s+)?", raw)
        parts = [p.strip().strip("'\"") for p in parts if p.strip()]
        parts = [p for p in parts if len(p) > 2 and "=" not in p]
        if parts:
            return parts

    return []


# ─── Dataclass ───────────────────────────────────────────────────────────────

@dataclass
class Criterion:
    criterion_id: str
    criterion_type: str
    weight: int
    description: str
    expected: str


# ─── Helpers ─────────────────────────────────────────────────────────────────

def _normalize_name(s: str) -> str:
    """Normalize name for fuzzy matching: lowercase, strip, underscores/hyphens to spaces, collapse whitespace."""
    s = s.lower().strip()
    s = s.replace("_", " ").replace("-", " ")
    return re.sub(r"\s+", " ", s)


def _to_str(val) -> str:
    """Convert expected field to string (may be list / float / int)."""
    if isinstance(val, str):
        return val
    if isinstance(val, list):
        return ", ".join(str(v) for v in val)
    return str(val)


def extract_file_meta(filepath: str) -> dict:
    """Extract structured file metadata."""
    meta: dict = {
        "filename": os.path.basename(filepath),
        "ext": os.path.splitext(filepath)[1].lower(),
        "size_bytes": os.path.getsize(filepath),
    }
    ext = meta["ext"]
    try:
        if ext in (".xlsx", ".xls", ".xlsm"):
            wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
            meta["sheet_names"] = wb.sheetnames
            meta["headers"] = {}
            meta["row_counts"] = {}
            for name in wb.sheetnames:
                ws = wb[name]
                rows = list(ws.iter_rows(max_row=1, values_only=True))
                meta["headers"][name] = [str(c) if c is not None else "" for c in rows[0]] if rows else []
                meta["row_counts"][name] = ws.max_row or 0
            wb.close()
        elif ext == ".pdf":
            with pdfplumber.open(filepath) as pdf:
                meta["page_count"] = len(pdf.pages)
        elif ext in (".docx", ".doc"):
            doc = DocxDocument(filepath)
            meta["headings"] = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
            meta["paragraph_count"] = len(doc.paragraphs)
            meta["table_count"] = len(doc.tables)
        elif ext in (".pptx", ".ppt"):
            prs = Presentation(filepath)
            meta["slide_count"] = len(prs.slides)
    except Exception as e:
        meta["parse_error"] = str(e)
    return meta


def try_rule_verify(
    criterion: Criterion,
    deliverables: list[str],
    all_meta: list[dict],
) -> dict | None:
    """Try to verify a criterion using rules. Returns result dict or None."""
    exp = criterion.expected

    # ── file existence ──
    filenames = parse_expected_filenames(exp)
    if filenames:
        found = {os.path.basename(f) for f in deliverables}
        missing = [fn for fn in filenames if fn not in found]
        if not missing:
            return {"satisfied": True, "evidence": f"All files found: {filenames}", "reason": "rule:file_exists"}
        else:
            return {"satisfied": False, "evidence": f"Missing files: {missing}", "reason": "rule:file_exists"}

    # ── sheet name existence ──
    sheets = parse_expected_sheets(exp)
    if sheets:
        all_sheets = set()
        for m in all_meta:
            all_sheets.update(_normalize_name(s) for s in m.get("sheet_names", []))
        missing = [s for s in sheets if _normalize_name(s) not in all_sheets]
        if not missing:
            return {"satisfied": True, "evidence": f"Sheets found: {sheets}", "reason": "rule:sheet_exists"}
        elif all_sheets:
            return {"satisfied": False, "evidence": f"Missing sheets: {missing}", "reason": "rule:sheet_exists"}

    # ── column name existence ──
    columns = parse_expected_columns(exp)
    if columns:
        all_headers = set()
        for m in all_meta:
            for hdrs in m.get("headers", {}).values():
                all_headers.update(_normalize_name(h) for h in hdrs if h)
        missing = [c for c in columns if _normalize_name(c) not in all_headers]
        if not missing:
            return {"satisfied": True, "evidence": f"All columns found", "reason": "rule:column_exists"}
        elif all_headers:
            return {"satisfied": False, "evidence": f"Missing columns: {missing}", "reason": "rule:column_exists"}

    # ── docx section existence ──
    sections = parse_expected_sections(exp)
    if sections:
        all_headings = set()
        for m in all_meta:
            all_headings.update(_normalize_name(h) for h in m.get("headings", []))
        if all_headings:
            missing_sections = [s for s in sections if _normalize_name(s) not in all_headings]
            if not missing_sections:
                return {"satisfied": True, "evidence": f"Sections found: {sections}", "reason": "rule:section_exists"}
            else:
                return {"satisfied": False, "evidence": f"Missing sections: {missing_sections}", "reason": "rule:section_exists"}

    return None


# ─── Standalone main ─────────────────────────────────────────────────────────

def _scan_deliverables(task_dir: str) -> list[str]:
    """Scan task_dir for deliverable file paths.

    Search priority:
      1. ``home/workspace`` / ``home`` / ``deliverable_files`` subdirs (legacy
         layouts where the agent runs in a nested home dir).
      2. ``task_dir`` root itself (current gdpval-synth layout — agents write
         deliverables straight into ``$TASK_WORKSPACE``). Subdirs holding
         input fixtures are skipped to avoid grading them as deliverables.
    """
    deliverables: list[str] = []
    extensions = (".xlsx", ".xls", ".xlsm", ".pdf", ".pptx", ".ppt", ".docx", ".doc")
    for subdir in ("home/workspace", "home", "deliverable_files"):
        candidate = os.path.join(task_dir, subdir)
        if os.path.isdir(candidate):
            for root, _, files in os.walk(candidate):
                for fname in files:
                    if fname.lower().endswith(extensions):
                        deliverables.append(os.path.join(root, fname))
            if deliverables:
                return deliverables

    # Fallback: scan task_dir root (current layout). Skip input/system dirs.
    skip_dirs = {"reference_files", "environment", "agent", "memory", "skills",
                 "deliverable_files", "home"}
    if os.path.isdir(task_dir):
        for entry in os.listdir(task_dir):
            full = os.path.join(task_dir, entry)
            if os.path.isfile(full) and entry.lower().endswith(extensions):
                deliverables.append(full)
            elif os.path.isdir(full) and entry not in skip_dirs:
                for root, _, files in os.walk(full):
                    for fname in files:
                        if fname.lower().endswith(extensions):
                            deliverables.append(os.path.join(root, fname))
    return deliverables


def main():
    """Standalone entry point. Reads env vars and outputs JudgerResult JSON."""
    task_workspace = os.environ.get("TASK_WORKSPACE", "")
    rubric_path = os.environ.get("RUBRIC_PATH", "")
    deliverable_spec_path = os.environ.get("DELIVERABLE_SPEC_PATH", "")
    judger_name = os.environ.get("JUDGER_NAME", "rule_judge")

    assert task_workspace, "TASK_WORKSPACE env var is required"
    assert rubric_path, "RUBRIC_PATH env var is required"

    # Load rubric
    with open(rubric_path) as f:
        rubric_data = json.load(f)

    raw_criteria = rubric_data.get("criteria", [])
    assert raw_criteria, f"No criteria found in rubric: {rubric_path}"

    # Filter rule criteria
    rule_criteria = [c for c in raw_criteria if c.get("judge_method") == "rule"]
    if not rule_criteria:
        # No rule criteria, output empty result
        result = {
            "judger_name": judger_name,
            "total": 0.0,
            "criteria": {},
            "metadata": {"summary": {"total": 0, "satisfied": 0, "not_satisfied": 0, "undetermined": 0}},
        }
        print(json.dumps(result, ensure_ascii=False))
        return

    # Scan deliverables
    deliverables = _scan_deliverables(task_workspace)

    # Extract file metadata
    all_meta: list[dict] = []
    for fpath in deliverables:
        try:
            all_meta.append(extract_file_meta(fpath))
        except Exception as e:
            all_meta.append({
                "filename": os.path.basename(fpath),
                "ext": os.path.splitext(fpath)[1],
                "parse_error": str(e),
            })

    # Run rule verification for each criterion
    criteria_results = []
    satisfied_count = 0
    not_satisfied_count = 0
    undetermined_count = 0

    for i, c in enumerate(rule_criteria):
        criterion = Criterion(
            criterion_id=c.get("criterion_id", f"c_{i}"),
            criterion_type=c.get("criterion_type", "unknown"),
            weight=c.get("weight", 1),
            description=_to_str(c.get("description", "")),
            expected=_to_str(c.get("expected", "")),
        )
        rule_result = try_rule_verify(criterion, deliverables, all_meta)

        if rule_result is not None:
            score = 1.0 if rule_result["satisfied"] else 0.0
            if rule_result["satisfied"]:
                satisfied_count += 1
            else:
                not_satisfied_count += 1
        else:
            score = -1.0  # undetermined, needs simple_judge
            undetermined_count += 1
            rule_result = {"satisfied": None, "evidence": "", "reason": "rule:no_match"}

        criteria_results.append({
            "criterion_id": criterion.criterion_id,
            "score": score,
            "satisfied": rule_result["satisfied"],
            "evidence": rule_result.get("evidence", ""),
            "reason": rule_result.get("reason", ""),
        })

    # Compute aggregate score: undetermined criteria are excluded from
    # the denominator (they belong to simple_judger, not us).
    scored = [r for r in criteria_results if r["score"] >= 0]
    total = (sum(r["score"] for r in scored) / len(scored)) if scored else 0.0

    criteria_map = {
        r["criterion_id"]: {"score": max(0.0, r["score"])}
        for r in criteria_results if r["score"] >= 0
    }

    result = {
        "judger_name": judger_name,
        "total": total,
        "criteria": criteria_map,
        "metadata": {
            "criteria_results": criteria_results,
            "summary": {
                "total": len(rule_criteria),
                "satisfied": satisfied_count,
                "not_satisfied": not_satisfied_count,
                "undetermined": undetermined_count,
            },
        },
    }
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
