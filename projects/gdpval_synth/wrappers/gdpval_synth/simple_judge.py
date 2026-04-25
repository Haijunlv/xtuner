"""LLM-based simple judger for gdpval_synth Agent Rollout.

Reads rubric.json, filters criteria with judge_method == "simple_judge",
reads agent output files, calls an LLM API to score each criterion,
and outputs a single JudgerResult JSON line to stdout.
"""

from __future__ import annotations

import json
import os
import re
import ssl
import sys
import urllib.request
from pathlib import Path

# ---------------------------------------------------------------------------
# Logging helper (all logs to stderr so stdout stays clean for JSON output)
# ---------------------------------------------------------------------------

def _log(msg: str) -> None:
    print(f"[simple_judge] {msg}", file=sys.stderr)


# ---------------------------------------------------------------------------
# File reading helpers
# ---------------------------------------------------------------------------

def _resolve_file(workspace: Path, spec_filename: str) -> Path | None:
    """Try full relative path first, then basename-only fallback."""
    candidate = workspace / spec_filename
    if candidate.exists():
        return candidate
    candidate = workspace / Path(spec_filename).name
    if candidate.exists():
        return candidate
    return None


def _read_xlsx(filepath: Path, max_rows: int = 20) -> str:
    """Read xlsx file, return text representation of first N rows per sheet."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    except Exception as exc:
        return f"[Error reading xlsx: {exc}]"

    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"=== Sheet: {sheet_name} ===")
        row_count = 0
        for row in ws.iter_rows(values_only=True):
            if row_count >= max_rows:
                parts.append(f"... (truncated after {max_rows} rows)")
                break
            cells = [str(c) if c is not None else "" for c in row]
            parts.append(" | ".join(cells))
            row_count += 1
        if row_count == 0:
            parts.append("(empty sheet)")
    wb.close()
    return "\n".join(parts)


def _read_docx(filepath: Path) -> str:
    """Read docx file, return paragraph text."""
    try:
        import docx
        doc = docx.Document(str(filepath))
    except Exception as exc:
        return f"[Error reading docx: {exc}]"

    parts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    return "\n".join(parts)


def _read_file_content(filepath: Path, max_chars: int = 3000) -> str:
    """Read file content based on extension, truncated to max_chars."""
    ext = filepath.suffix.lower()

    if ext == ".xlsx":
        content = _read_xlsx(filepath)
    elif ext == ".docx":
        content = _read_docx(filepath)
    elif ext in (".csv", ".txt", ".json", ".md", ".py", ".log"):
        try:
            content = filepath.read_text(encoding="utf-8", errors="replace")
        except Exception as exc:
            content = f"[Error reading {filepath.name}: {exc}]"
    else:
        try:
            content = filepath.read_text(encoding="utf-8", errors="replace")
        except Exception:
            content = f"[Binary or unreadable file: {filepath.name}]"

    if len(content) > max_chars:
        content = content[:max_chars] + f"\n... (truncated at {max_chars} chars)"
    return content


# ---------------------------------------------------------------------------
# LLM API call
# ---------------------------------------------------------------------------

def _call_llm(
    base_url: str,
    api_key: str,
    model: str,
    prompt: str,
    timeout: int = 120,
) -> str:
    """Call OpenAI-compatible /chat/completions endpoint via urllib."""
    url = base_url.rstrip("/") + "/chat/completions"
    payload = json.dumps({
        "model": model,
        "messages": [
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.1,
        "max_tokens": 512,
    }).encode("utf-8")

    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {api_key}",
    }

    req = urllib.request.Request(url, data=payload, headers=headers, method="POST")

    # Allow self-signed certs in internal environments
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE

    with urllib.request.urlopen(req, timeout=timeout, context=ctx) as resp:
        body = json.loads(resp.read().decode("utf-8"))

    return body["choices"][0]["message"]["content"]


def _parse_score(text: str) -> float:
    """Extract score from LLM response. Look for 'Score: N.N' pattern."""
    # Pattern 1: explicit "Score: X.X"
    m = re.search(r"[Ss]core\s*[:=]\s*([01](?:\.\d+)?)", text)
    if m:
        return float(m.group(1))

    # Pattern 2: standalone 0-1 float on its own line
    for line in reversed(text.strip().splitlines()):
        line = line.strip()
        m2 = re.match(r"^([01](?:\.\d+)?)$", line)
        if m2:
            return float(m2.group(1))

    # Pattern 3: any 0-1 float in the text (last occurrence)
    all_floats = re.findall(r"\b([01]\.\d+)\b", text)
    if all_floats:
        return float(all_floats[-1])

    _log(f"Could not parse score from LLM response: {text[:200]}")
    return 0.0


# ---------------------------------------------------------------------------
# Build LLM prompt
# ---------------------------------------------------------------------------

def _build_prompt(
    criterion: dict,
    file_contents: dict[str, str],
) -> str:
    """Build a scoring prompt for the LLM."""
    description = criterion.get("description", "")
    expected = criterion.get("expected", "")
    criterion_type = criterion.get("criterion_type", "")

    files_section = ""
    for fname, content in file_contents.items():
        files_section += f"\n--- File: {fname} ---\n{content}\n"

    prompt = f"""You are an expert evaluator. Score how well the agent's output satisfies the following criterion.

## Criterion
- Type: {criterion_type}
- Description: {description}
- Expected: {expected}

## Agent Output Files
{files_section if files_section else "(No output files found)"}

## Instructions
Evaluate the agent's output against the criterion above.
- Score 1.0 if the criterion is fully satisfied.
- Score 0.0 if the criterion is not satisfied at all.
- Use intermediate values (e.g., 0.5) for partial satisfaction.

Provide a brief explanation, then state your score on the last line in this exact format:
Score: X.X
"""
    return prompt


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main() -> None:
    workspace_str = os.environ.get("TASK_WORKSPACE", "")
    rubric_path_str = os.environ.get("RUBRIC_PATH", "")
    spec_path_str = os.environ.get("DELIVERABLE_SPEC_PATH", "")
    judger_name = os.environ.get("JUDGER_NAME", "simple_judger")

    llm_base_url = os.environ.get("RL_LLM_BASE_URL", "")
    llm_api_key = os.environ.get("RL_LLM_API_KEY", "")
    llm_model = os.environ.get("RL_LLM_MODEL", "")

    if not workspace_str:
        print(json.dumps({"judger_name": judger_name, "total": 0.0, "error": "TASK_WORKSPACE not set"}))
        return

    workspace = Path(workspace_str)

    # Check LLM config
    if not llm_base_url:
        print(json.dumps({
            "judger_name": judger_name,
            "total": 0.0,
            "criteria": {},
            "error": "RL_LLM_BASE_URL not set, cannot run LLM-based judging",
            "metadata": {"simple_criteria_count": 0, "llm_model": ""},
        }))
        return

    # Read rubric
    rubric_path = Path(rubric_path_str) if rubric_path_str else workspace / "rubric.json"
    try:
        rubric = json.loads(rubric_path.read_text(encoding="utf-8"))
    except Exception as exc:
        print(json.dumps({"judger_name": judger_name, "total": 0.0, "error": f"Cannot read rubric: {exc}"}))
        return

    # Read deliverable spec
    spec_path = Path(spec_path_str) if spec_path_str else workspace / "deliverable_spec.json"
    try:
        spec = json.loads(spec_path.read_text(encoding="utf-8"))
    except Exception as exc:
        _log(f"Cannot read deliverable_spec: {exc}, using empty spec")
        spec = []

    # Filter simple_judge criteria
    all_criteria = rubric.get("criteria", [])
    simple_criteria = [c for c in all_criteria if c.get("judge_method") == "simple_judge"]

    if not simple_criteria:
        print(json.dumps({
            "judger_name": judger_name,
            "total": 1.0,
            "criteria": {},
            "metadata": {"skipped": True, "simple_criteria_count": 0, "llm_model": llm_model},
        }))
        return

    _log(f"Found {len(simple_criteria)} simple_judge criteria out of {len(all_criteria)} total")

    # Pre-read all deliverable files
    file_contents: dict[str, str] = {}
    for entry in spec:
        fn = entry.get("filename", "")
        resolved = _resolve_file(workspace, fn)
        if resolved is not None:
            file_contents[fn] = _read_file_content(resolved)
        else:
            _log(f"  deliverable file not found: {fn}")

    # Judge each criterion
    criteria_results: dict[str, dict] = {}

    for criterion in simple_criteria:
        cid = criterion.get("criterion_id", "unknown")
        weight = float(criterion.get("weight", 1.0))

        prompt = _build_prompt(criterion, file_contents)

        try:
            response_text = _call_llm(llm_base_url, llm_api_key, llm_model, prompt)
            score = _parse_score(response_text)
            _log(f"  {cid}: score={score:.4f}, weight={weight}")
        except Exception as exc:
            _log(f"  {cid}: LLM call failed: {exc}")
            score = 0.0

        criteria_results[cid] = {"score": round(score, 4), "weight": weight}

    # Compute weighted total
    total_weight = sum(r["weight"] for r in criteria_results.values())
    if total_weight > 0:
        total_score = sum(r["score"] * r["weight"] for r in criteria_results.values()) / total_weight
    else:
        total_score = 0.0

    result = {
        "judger_name": judger_name,
        "total": round(total_score, 4),
        "criteria": criteria_results,
        "metadata": {
            "simple_criteria_count": len(simple_criteria),
            "llm_model": llm_model,
        },
    }

    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
